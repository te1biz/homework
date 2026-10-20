
import csv
import json
import logging
from pathlib import Path
from typing import Any

from src.processing import filter_by_state, sort_by_date
from src.search import process_bank_search
from src.utils import load_transactions
from src.widget import get_date, mask_account_card


logger = logging.getLogger(__name__)

AVAILABLE_STATUSES = {"EXECUTED", "CANCELED", "PENDING"}


def load_csv_transactions(file_path: str | Path) -> list[dict[str, Any]]:
    """Загружает транзакции из CSV-файла."""
    try:
        with open(file_path, encoding="utf-8-sig", newline="") as file:
            return list(csv.DictReader(file))
    except (OSError, csv.Error) as error:
        logger.error("Не удалось загрузить CSV: %s", error)
        return []


def load_xlsx_transactions(file_path: str | Path) -> list[dict[str, Any]]:
    """Загружает транзакции из XLSX-файла."""
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        worksheet = workbook.active

        if worksheet is None:
            workbook.close()
            return []

        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            workbook.close()
            return []

        column_names = [str(value) if value is not None else "" for value in headers]
        transactions: list[dict[str, Any]] = []

        for row in rows:
            transaction = {
                column_names[index]: value
                for index, value in enumerate(row)
                if index < len(column_names) and column_names[index]
            }
            transactions.append(transaction)

        workbook.close()
        return transactions

    except (OSError, ValueError, ImportError) as error:
        logger.error("Не удалось загрузить XLSX: %s", error)
        return []


def get_file_type() -> str | None:
    """Запрашивает у пользователя формат файла."""
    print(
        "Привет! Добро пожаловать в программу работы "
        "с банковскими транзакциями.\n"
        "Выберите необходимый пункт меню:\n"
        "1. Получить информацию о транзакциях из JSON-файла\n"
        "2. Получить информацию о транзакциях из CSV-файла\n"
        "3. Получить информацию о транзакциях из XLSX-файла"
    )

    choices = {
        "1": "JSON",
        "2": "CSV",
        "3": "XLSX",
    }

    while True:
        choice = input().strip()

        if choice in choices:
            file_type = choices[choice]
            print(f"Для обработки выбран {file_type}-файл.")
            return file_type

        print("Неверный пункт меню. Введите 1, 2 или 3.")


def get_transactions(file_type: str) -> list[dict[str, Any]]:
    """Загружает транзакции из выбранного формата."""
    default_paths = {
        "JSON": Path("data/operations.json"),
        "CSV": Path("data/transactions.csv"),
        "XLSX": Path("data/transactions.xlsx"),
    }

    file_path = default_paths[file_type]

    if not file_path.exists():
        entered_path = input(
            f"Файл {file_path} не найден. Введите путь к файлу: "
        ).strip()

        if not entered_path:
            print("Путь к файлу не указан.")
            return []

        file_path = Path(entered_path)

    if file_type == "JSON":
        return load_transactions(file_path)

    if file_type == "CSV":
        return load_csv_transactions(file_path)

    return load_xlsx_transactions(file_path)


def get_status() -> str:
    """Запрашивает допустимый статус транзакций."""
    while True:
        print(
            "Введите статус, по которому необходимо выполнить фильтрацию.\n"
            "Доступные для фильтровки статусы: EXECUTED, CANCELED, PENDING"
        )
        status = input().strip().upper()

        if status in AVAILABLE_STATUSES:
            print(f'Операции отфильтрованы по статусу "{status}"')
            return status

        print(f'Статус операции "{status}" недоступен.')


def ask_yes_no(question: str) -> bool:
    """Запрашивает ответ да или нет."""
    while True:
        answer = input(f"{question} Да/Нет: ").strip().lower()

        if answer in {"да", "д", "yes", "y"}:
            return True

        if answer in {"нет", "н", "no", "n"}:
            return False

        print("Пожалуйста, ответьте Да или Нет.")


def sort_transactions(
    transactions: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """При необходимости сортирует транзакции по дате."""
    if not ask_yes_no("Отсортировать операции по дате?"):
        return transactions

    while True:
        direction = input(
            "Отсортировать по возрастанию или по убыванию? "
        ).strip().lower()

        if direction in {"по возрастанию", "возрастание", "возрастанию"}:
            return sort_by_date(transactions, is_reversed=False)

        if direction in {"по убыванию", "убывание", "убыванию"}:
            return sort_by_date(transactions, is_reversed=True)

        print("Введите «по возрастанию» или «по убыванию».")


def filter_ruble_transactions(
    transactions: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Оставляет только рублёвые транзакции, если это выбрано."""
    if not ask_yes_no("Выводить только рублевые транзакции?"):
        return transactions

    result: list[dict[str, Any]] = []

    for transaction in transactions:
        operation_amount = transaction.get("operationAmount", {})
        currency = operation_amount.get("currency", {})

        if isinstance(currency, dict) and currency.get("code") == "RUB":
            result.append(transaction)
            continue

        # Поддержка строкового формата из CSV/XLSX.
        if transaction.get("currency_code") == "RUB":
            result.append(transaction)

    return result


def search_transactions(
    transactions: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """При необходимости фильтрует операции по описанию."""
    if not ask_yes_no(
        "Отфильтровать список транзакций по определенному слову в описании?"
    ):
        return transactions

    search = input("Введите слово для поиска: ").strip()

    if not search:
        return transactions

    return process_bank_search(transactions, search)


def get_transaction_amount(
    transaction: dict[str, Any],
) -> tuple[str, str]:
    """Возвращает сумму и код валюты транзакции."""
    operation_amount = transaction.get("operationAmount", {})

    if isinstance(operation_amount, dict):
        amount = operation_amount.get("amount", "0")
        currency = operation_amount.get("currency", {})

        if isinstance(currency, dict):
            currency_code = str(currency.get("code", ""))
        else:
            currency_code = ""

        return str(amount), currency_code

    # Поддержка таблиц, в которых сумма и валюта находятся в отдельных колонках.
    amount = transaction.get("amount", "0")
    currency_code = transaction.get("currency_code", "")
    return str(amount), str(currency_code)


def print_transactions(transactions: list[dict[str, Any]]) -> None:
    """Выводит список транзакций в читаемом виде."""
    print("Распечатываю итоговый список транзакций...")

    if not transactions:
        print("Не найдено ни одной транзакции, подходящей под ваши условия фильтрации")
        return

    print(f"\nВсего банковских операций в выборке: {len(transactions)}\n")

    for transaction in transactions:
        date_value = transaction.get("date", "")
        description = transaction.get("description", "Без описания")

        try:
            formatted_date = get_date(str(date_value))
        except (ValueError, TypeError):
            formatted_date = str(date_value)

        print(f"{formatted_date} {description}")

        from_account = transaction.get("from")
        to_account = transaction.get("to")

        if from_account:
            try:
                from_account = mask_account_card(str(from_account))
            except ValueError:
                from_account = str(from_account)

        if to_account:
            try:
                to_account = mask_account_card(str(to_account))
            except ValueError:
                to_account = str(to_account)

        if from_account and to_account:
            print(f"{from_account} -> {to_account}")
        elif to_account:
            print(to_account)
        elif from_account:
            print(from_account)

        amount, currency = get_transaction_amount(transaction)
        print(f"Сумма: {amount} {currency}".strip())
        print()


def main() -> None:
    """Запускает программу обработки банковских транзакций."""
    file_type = get_file_type()

    if file_type is None:
        return

    transactions = get_transactions(file_type)

    if not transactions:
        print("Не удалось загрузить транзакции или файл пуст.")
        return

    status = get_status()
    filtered = filter_by_state(transactions, status)

    filtered = sort_transactions(filtered)
    filtered = filter_ruble_transactions(filtered)
    filtered = search_transactions(filtered)

    print_transactions(filtered)


if __name__ == "__main__":
    main()
